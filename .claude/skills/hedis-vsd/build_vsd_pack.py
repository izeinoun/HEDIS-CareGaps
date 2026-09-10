#!/usr/bin/env python3
"""Build a year-scoped value-set pack from the NCQA HEDIS Value Set Directory (VSD).

The VSD is the licensed NCQA workbook that maps each value-set NAME to its exhaustive
code list (ICD-10, CPT, HCPCS, LOINC, SNOMED, CVX…) for one measurement year. The rules
pack (from `hedis-spec-to-rules`) deliberately carries value sets by NAME only; this
builder produces the companion code dictionary the analysis skill and the reviewer use
to resolve names → codes and to validate a captured code's membership.

Deterministic — no LLM. It parses a normalized CSV or .xlsx export of the directory
(tolerant column mapping) and groups rows by value-set name. Codes are transcribed
verbatim; nothing is inferred.

Guardrails: the VSD is licensed content — the pack is an INTERNAL artifact, kept
separate from the (shareable) rules pack, never redistributed, and one-per-year (codes
are not stable across years).

Usage:
  python build_vsd_pack.py --input VSD.csv --year MY2027 [--title "..."] [--out pack.json]
"""
import argparse
import csv
import json
import re
import sys

VSD_VERSION = "1.0"

# Tolerant header mapping — accept the common variants a VSD export uses.
_COLS = {
    "value_set": ("value set name", "value_set_name", "value set", "valueset", "vs name"),
    "oid":       ("oid", "value set oid", "code system oid"),
    "system":    ("code system", "code_system", "system", "codesystem"),
    "code":      ("code", "concept code"),
    "desc":      ("description", "definition", "display", "code description", "concept name"),
}


def _norm(s):
    return re.sub(r"\s+", " ", (s or "")).strip()


def _map_headers(header):
    idx = {}
    low = [(_norm(h).lower()) for h in header]
    for field, names in _COLS.items():
        for i, h in enumerate(low):
            if h in names:
                idx[field] = i
                break
    return idx


def _rows_from_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as fh:
        reader = csv.reader(fh)
        rows = list(reader)
    return rows


def _rows_from_xlsx(path):
    from openpyxl import load_workbook  # optional dependency
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    return [[("" if c is None else str(c)) for c in row] for row in ws.iter_rows(values_only=True)]


def build(rows, measurement_year, directory_title=None, publisher="NCQA",
          effective=None, loaded_at=None):
    """Group VSD rows into a value-set pack. `rows` is a list of lists, first row = header."""
    if not rows:
        raise ValueError("no rows")
    idx = _map_headers(rows[0])
    for req in ("value_set", "system", "code"):
        if req not in idx:
            raise ValueError(f"could not find a '{req}' column in the header {rows[0]!r}")

    value_sets = {}
    for r in rows[1:]:
        def cell(field):
            i = idx.get(field)
            return _norm(r[i]) if i is not None and i < len(r) else ""
        name = cell("value_set")
        code = cell("code")
        if not name or not code:
            continue
        vs = value_sets.setdefault(name, {"oid": cell("oid") or None, "codes": []})
        if not vs["oid"] and cell("oid"):
            vs["oid"] = cell("oid")
        vs["codes"].append({"system": cell("system"), "code": code, "description": cell("desc") or None})

    yr = re.search(r"(\d{4})", measurement_year or "")
    eff = effective or ({"start": f"{yr.group(1)}-01-01", "end": f"{yr.group(1)}-12-31"} if yr else None)
    return {
        "vsd_version": VSD_VERSION,
        "source": {
            "publisher": publisher,
            "directory_title": directory_title or f"HEDIS {measurement_year} Value Set Directory",
            "measurement_year": measurement_year,
            "effective": eff,
            "loaded_at": loaded_at,
            "license": "Licensed NCQA Value Set Directory content — internal use only, not for "
                       "redistribution. Kept separate from the shareable rules pack.",
        },
        "value_sets": value_sets,
    }


def validate(pack):
    """Structural + sanity checks; returns a list of 'ERROR:'/'WARN:' lines (empty = clean)."""
    out = []
    src = pack.get("source") or {}
    if not src.get("measurement_year"):
        out.append("ERROR: source.measurement_year is required (codes are not stable across years)")
    vs = pack.get("value_sets")
    if not isinstance(vs, dict) or not vs:
        out.append("ERROR: value_sets must be a non-empty object")
        return out
    for name, body in vs.items():
        codes = (body or {}).get("codes") or []
        if not codes:
            out.append(f"WARN: value set '{name}' has no codes")
        for c in codes:
            if not c.get("system") or not c.get("code"):
                out.append(f"ERROR: value set '{name}' has a code missing system/code")
                break
    return out


def main():
    ap = argparse.ArgumentParser(description="Build a value-set pack from an NCQA VSD export.")
    ap.add_argument("--input", required=True, help="VSD export (.csv or .xlsx)")
    ap.add_argument("--year", required=True, help="measurement year, e.g. MY2027")
    ap.add_argument("--title")
    ap.add_argument("--loaded-at", help="ISO date the directory was loaded (pass in; not guessed)")
    ap.add_argument("--out")
    args = ap.parse_args()
    rows = _rows_from_xlsx(args.input) if args.input.lower().endswith(".xlsx") else _rows_from_csv(args.input)
    pack = build(rows, args.year, directory_title=args.title, loaded_at=args.loaded_at)
    problems = validate(pack)
    for line in problems:
        print(line, file=sys.stderr)
    n_vs = len(pack["value_sets"])
    n_codes = sum(len(v["codes"]) for v in pack["value_sets"].values())
    out_json = json.dumps(pack, indent=2)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as fh:
            fh.write(out_json)
        print(f"Wrote {args.out}: {n_vs} value sets, {n_codes} codes ({args.year}).")
    else:
        print(out_json)
    return 1 if any(p.startswith("ERROR") for p in problems) else 0


if __name__ == "__main__":
    sys.exit(main())
